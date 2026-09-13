"""Generate the ResolveAI v5 judge deck without any external assets."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUT=Path("submission/ResolveAI_V5_Presentation.pptx")
NAVY=RGBColor(7,19,31);PANEL=RGBColor(13,32,51);WHITE=RGBColor(238,246,253);MUTED=RGBColor(150,180,204);MINT=RGBColor(56,214,176);BLUE=RGBColor(84,169,255);AMBER=RGBColor(247,185,85);RED=RGBColor(255,107,115)

def box(slide,x,y,w,h,color=PANEL,radius=True):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));shape.fill.solid();shape.fill.fore_color.rgb=color;shape.line.color.rgb=color;return shape
def text(slide,content,x,y,w,h,size=18,color=WHITE,bold=False,align=None):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.clear();tf.word_wrap=True;p=tf.paragraphs[0];p.text=content;p.font.name="Aptos";p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color
    if align:p.alignment=align
    return tb
def title(slide,kicker,headline,sub=""):
    text(slide,kicker.upper(),.7,.42,11,.28,10,MINT,True);text(slide,headline,.7,.72,12,0.75,28,WHITE,True);text(slide,sub,.72,1.52,11.7,.42,12,MUTED)
def card(slide,heading,body,x,y,w,h,accent=MINT):
    box(slide,x,y,w,h);box(slide,x,y,.07,h,accent,False);text(slide,heading,x+.25,y+.18,w-.4,.3,14,WHITE,True);text(slide,body,x+.25,y+.6,w-.45,h-.7,11,MUTED)
def base(prs):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=NAVY;return slide

prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
# 1
s=base(prs);text(s,"RESOLVEAI",.72,.72,3,.3,13,MINT,True);text(s,"Autonomous resolution,\nvisibly grounded.",.7,1.25,8,1.45,36,WHITE,True);text(s,"An agentic customer-operations system that investigates, acts through controlled tools, adapts to real constraints, and verifies outcomes.",.75,3.02,7.2,.7,16,MUTED);box(s,8.7,1.05,3.8,4.8);text(s,"OBSERVE",9.1,1.5,2.8,.3,14,BLUE,True);text(s,"DECIDE",9.1,2.45,2.8,.3,14,MINT,True);text(s,"ACT",9.1,3.4,2.8,.3,14,AMBER,True);text(s,"VERIFY",9.1,4.35,2.8,.3,14,WHITE,True);text(s,"Tech Zephyr 4.0 · Smart Automation",.75,6.68,5,.25,11,MUTED)
# 2
s=base(prs);title(s,"The problem","Support replies are not customer outcomes.","Traditional chatbots can answer—but cannot safely resolve an operational case.");card(s,"No enterprise truth","Replies often assume order, payment, and inventory facts instead of reading them.",.72,2.3,3.8,2.25,RED);card(s,"No controlled action","A chat response cannot create a refund, cancel an order, or reserve stock safely.",4.77,2.3,3.8,2.25,AMBER);card(s,"No proof of effect","Even a successful action claim is incomplete without an independent state check.",8.82,2.3,3.8,2.25,BLUE)
# 3
s=base(prs);title(s,"The product","A bounded agent loop—not a scripted issue router.","The model selects one tool at a time. Python validates, executes, and enforces policy.");steps=[("1","Observe","Customer request + authorized case snapshot + prior tool transcript"),("2","Ground","TF-IDF Policy RAG + computed policy gates + Bitext intent hint"),("3","Act","Allow-listed, idempotent enterprise tools mutate only when permitted"),("4","Verify","Independent DB reads confirm the exact requested state transition")]
for i,(n,h,b) in enumerate(steps):
 x=.72+i*3.13;box(s,x,2.4,2.75,2.65);text(s,n,x+.2,2.62,.35,.35,23,MINT,True);text(s,h,x+.2,3.12,2.3,.3,15,WHITE,True);text(s,b,x+.2,3.62,2.3,1,11,MUTED)
# 4
s=base(prs);title(s,"Safety architecture","Policy interpretation informs. Deterministic rules authorize.","No LLM output and no seed-data boolean can bypass the service gate.");card(s,"Computed policy rules","Delivery date → 30-day refund/replacement window\nOrder & payment state → eligibility\nShipment status → cancellation gate\nAmount > ₹20,000 → human approval",.72,2.25,4,3.1,MINT);card(s,"Policy RAG","Ranked policy snippets appear in the agent context and UI. Each deterministic gate reports a matching policy reference such as refund-window.",4.72,2.25,3.85,3.1,BLUE);card(s,"Execution controls","Server-side re-checks\nIdempotency keys\nMutation cap\nTransient retry once\nCircuit breaker\nAudit event in the same transaction",8.57,2.25,4.03,3.1,AMBER)
# 5
s=base(prs);title(s,"Generalization & identity","Natural language is a hint; authorization is never inferred.","The system remains safe when phrasing is novel, incomplete, or ambiguous.");card(s,"Bitext intent hint","A committed cache from Bitext’s 26,872 customer-support examples feeds a local TF-IDF classifier. It improves natural-language entity resolution without runtime network calls.",.72,2.25,3.85,3.05,MINT);card(s,"Exact IDs win","CASE-### and ORD-### resolution stays authoritative. Intent only narrows candidates after the exact-ID path fails.",4.74,2.25,3.85,3.05,BLUE);card(s,"No cross-customer leak","Unscoped ambiguity returns NEEDS CLARIFICATION. API routes reject a case/customer-context ownership mismatch with 403.",8.76,2.25,3.85,3.05,RED)
# 6
s=base(prs);title(s,"Live demo: adaptation","A customer asks for a replacement. The system discovers it cannot fulfil one.","Use CASE-100: “My headphones arrived damaged and I want a replacement.”");flow=[("Policy RAG","Retrieves damage, replacement and refund policy evidence",BLUE),("Inventory","Zero available headphones across warehouses",RED),("Replan","Replacement is blocked; evaluate refund",AMBER),("Refund + verify","Payment ledger transitions and independent check passes",MINT)]
for i,(h,b,c) in enumerate(flow):
 y=2.2+i*.86;box(s,1.0,y,11.2,.63);text(s,h,1.28,y+.13,2.2,.25,13,c,True);text(s,b,3.38,y+.14,8.2,.25,11,MUTED)
# 7
s=base(prs);title(s,"Live evidence surfaces","The demo makes every claim inspectable.","React SSE trace is the presentation view; Streamlit remains the operations console.");card(s,"Agent trace","Turn-by-turn reasoning, tool selected, structured result, adaptation, and terminal outcome arrive over Server-Sent Events.",.72,2.25,3.85,2.8,MINT);card(s,"Policy evidence","Retrieved policy snippets and similarity scores are shown beside the trace—RAG is observable, not decorative.",4.74,2.25,3.85,2.8,BLUE);card(s,"State & audit","Authoritative order/payment/inventory snapshot plus persistent before → after audit rows and independent verification.",8.76,2.25,3.85,2.8,AMBER)
# 8
s=base(prs);title(s,"Proof, resilience & deployment","Built to show exactly what happened when a tool or provider fails.","A visible fallback is safer than a silent downgrade.");card(s,"Provider resilience","Gemini and OpenRouter receive one retry. Quota-limited providers cool down; offline test stub is visibly labeled—not presented as live reasoning.",.72,2.25,3.85,2.9,RED);card(s,"Test evidence","27 passing tests cover golden flows, policy date math, partial refund, policy retrieval, Bitext paraphrases, cross-customer isolation, and API ownership checks.",4.74,2.25,3.85,2.9,MINT);card(s,"Run anywhere","FastAPI API · React SSE demo · Streamlit ops view · Docker Compose · GitHub Actions CI · synthetic SQLite system of record.",8.76,2.25,3.85,2.9,BLUE)
# 9
s=base(prs);title(s,"Why ResolveAI","It resolves outcomes—and proves them.","The differentiator is not a fluent answer. It is a controlled, verified operational loop.");text(s,"Investigate real state  →  choose a safe tool  →  adapt to constraints  →  verify the database effect",.9,2.58,11.6,.55,20,MINT,True,PP_ALIGN.CENTER);text(s,"Demo URLs: React live trace :5173 · Streamlit ops :8501 · API health :8000/healthz",.9,4.25,11.6,.35,13,MUTED,False,PP_ALIGN.CENTER);text(s,"Thank you",.9,5.6,11.6,.5,25,WHITE,True,PP_ALIGN.CENTER)

OUT.parent.mkdir(exist_ok=True);prs.save(OUT);print(OUT)
